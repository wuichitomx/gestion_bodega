import io
import math
import os
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook
from vista_regional import (ADITIVOS, catalogo, codigo, datos_mapa, estado_seleccionado,
                           filtrar_estado, grafica_kpi, leer_reporte, numero, referencia_regional)


def libro(filas, total=True):
    w = Workbook()
    s = w.active
    s.append(["Extracto de prueba"])
    cab = ["No. Linea", "Código de Almacén", "Nombre de Almacén", *ADITIVOS, "Mt2"]
    s.append(cab)
    for i, fila in enumerate(filas, 1):
        s.append([i, fila.get("codigo", f"T{i}"), "Tienda de prueba", *[fila.get(k, 0) for k in ADITIVOS], fila.get("Mt2", 10)])
    if total:
        s.append(["TOTALES", None, None, *[sum(f.get(k, 0) for f in filas) for k in ADITIVOS], sum(f.get("Mt2", 10) or 0 for f in filas)])
    b = io.BytesIO()
    w.save(b)
    return b.getvalue()


class LectorTests(unittest.TestCase):
    def test_numeros(self):
        for v in (1234.56, "$1,234.56", "1.234,56", "1 234,56"):
            self.assertEqual(numero(v), 1234.56)
        self.assertEqual(numero("(1,234.56)"), -1234.56)
        self.assertEqual(numero("1,234"), 1234)
        for v in (None, True, "abc", float("inf"), "NaN"):
            with self.assertRaises(ValueError):
                numero(v)
        self.assertTrue(math.isnan(numero(None, opcional=True)))
        self.assertEqual(codigo(" z1gex "), "Z1GEX")
        self.assertEqual(codigo(123.0), "123")
        self.assertEqual(codigo("00123"), "00123")

    def test_totales_ponderacion_y_area(self):
        d, control, _, _ = leer_reporte(libro([
            {"Venta": 100, "Unidades": 2, "# Doc": 1, "Venta Bruta": 200, "Descuento": 100},
            {"Venta": 900, "Unidades": 18, "# Doc": 3, "Venta Bruta": 1000, "Descuento": 100, "Mt2": 0}]))
        self.assertEqual(len(d), 2)
        self.assertTrue(control.Coincide.all())
        self.assertTrue(math.isnan(d.iloc[1]["Venta x Mt2"]))
        r = referencia_regional(d)
        self.assertEqual(r["Venta"], 500)
        self.assertEqual(r["UPT"], 5)
        self.assertEqual(r["Venta x Mt2"], 10)
        self.assertAlmostEqual(r["% Dcto"], 1/6)

    def test_area_y_denominadores_invalidos(self):
        for area in (0, -1, None):
            d, _, _, _ = leer_reporte(libro([{"Mt2": area}], total=False))
            for k in ("Mt2", "Venta x Mt2", "UPT", "ATV", "ASP", "% Dcto"):
                self.assertTrue(math.isnan(d.iloc[0][k]), k)
            self.assertTrue(math.isnan(referencia_regional(d)["Venta x Mt2"]))

    def test_duplicado_rechazado(self):
        with self.assertRaisesRegex(ValueError, "duplicado"):
            leer_reporte(libro([{"codigo": "x"}, {"codigo": " X "}]))

    def test_graficas_referencia_completa_y_nd(self):
        d, _, _, _ = leer_reporte(libro([
            {"codigo": "Z1GEX", "Venta": 100, "Unidades": 2, "# Doc": 1},
            {"codigo": "Z1GPP", "Venta": 900, "Unidades": 18, "# Doc": 3, "Mt2": 0}]))
        ref = referencia_regional(d)
        self.assertEqual(ref["ATV"], 250)
        self.assertEqual(ref["ASP"], 50)
        for k in ("UPT", "ASP", "ATV", "Venta x Mt2"):
            spec = grafica_kpi(filtrar_estado(d, "Oaxaca"), k, ref[k]).to_dict()
            reglas = [v for rows in spec["datasets"].values() for v in rows if "referencia" in v]
            self.assertEqual(reglas, [{"referencia": ref[k]}])
            filas = [v for rows in spec["datasets"].values() for v in rows if "etiqueta" in v]
            self.assertEqual([v["codigo"] for v in filas], [] if k == "Venta x Mt2" else ["Z1GPP"])
        sin_area = d[d.codigo == "Z1GPP"]
        spec = grafica_kpi(sin_area, "Venta x Mt2", referencia_regional(sin_area)["Venta x Mt2"]).to_dict()
        self.assertFalse(any("referencia" in v for rows in spec["datasets"].values() for v in rows))

    def test_ausencia_totales_y_catalogo(self):
        d, c, avisos, _ = leer_reporte(libro([{"codigo": "desconocido"}], total=False))
        self.assertTrue(c.empty)
        self.assertIn("No hay fila TOTALES", avisos[0])
        self.assertEqual(d.iloc[0].estado, "Ubicación pendiente")
        self.assertEqual(len(catalogo()), 17)
        self.assertFalse(catalogo().cve_ent.isna().any())
        ubicaciones = catalogo().set_index("codigo")
        for cod in ("Z1CFF", "Z1IQF"):
            self.assertEqual(ubicaciones.loc[cod, "estado"], "Michoacán")
        self.assertEqual(ubicaciones.loc["Z1GKM", "estado"], "Jalisco")
        self.assertEqual(ubicaciones.loc["Z1GKP", "estado"], "Estado de México")

    def test_diferencia_totales_visible(self):
        w = load_workbook(io.BytesIO(libro([{"Venta": 100}])))
        w.active.cell(4, 9, 999)
        b = io.BytesIO(); w.save(b)
        _, c, a, _ = leer_reporte(b.getvalue())
        self.assertFalse(c.Coincide.all())
        self.assertIn("diferencias", a[0])

    def test_navegacion_estado(self):
        d, _, _, _ = leer_reporte(libro([{"codigo": "Z1GEX"}, {"codigo": "DESCONOCIDO"}]))
        self.assertEqual(len(filtrar_estado(d, "Toda la región")), 2)
        self.assertEqual(filtrar_estado(d, "Michoacán").iloc[0].codigo, "Z1GEX")
        self.assertEqual(len(filtrar_estado(d, "Ubicación pendiente")), 1)
        self.assertTrue(filtrar_estado(d, "Sonora").empty)
        features = datos_mapa(d)
        self.assertEqual(len(features), 32)
        for f in features:
            for poligono in f["geometry"]["coordinates"]:
                for i, anillo in enumerate(poligono):
                    self.assertEqual(anillo[0], anillo[-1])
                    area = sum(a[0]*b[1]-b[0]*a[1] for a, b in zip(anillo, anillo[1:]))
                    self.assertLessEqual(area if i == 0 else -area, 0)
        self.assertEqual(estado_seleccionado({"selection": {"estado": [{"cve_ent": "16"}]}}, features), "Michoacán")
        self.assertEqual(estado_seleccionado({"selection": {"estado": {"cve_ent": ["31"]}}}, features), "Yucatán")
        self.assertIsNone(estado_seleccionado({}, features))

    @unittest.skipUnless(os.environ.get("REGIONAL_EXCEL_REAL"), "Define REGIONAL_EXCEL_REAL para probar el archivo privado sin copiarlo al repo")
    def test_excel_real(self):
        p = Path(os.environ["REGIONAL_EXCEL_REAL"])
        d, c, avisos, _ = leer_reporte(p)
        self.assertEqual(len(d), 17)
        self.assertEqual(set(d.codigo), set(catalogo().codigo))
        self.assertTrue(c.Coincide.all())
        self.assertEqual(avisos, [])
        self.assertAlmostEqual(d.Venta.sum(), 13674839.30)
        self.assertAlmostEqual(d["Venta Neta"].sum(), 11788656.39)
        self.assertEqual(d.Unidades.sum(), 10124)
        self.assertEqual(d["# Doc"].sum(), 6477)
        self.assertAlmostEqual(referencia_regional(d)["Venta x Mt2"], 3414.77412693906)
        self.assertTrue(math.isnan(d.set_index("codigo").loc["Z1GPP", "Venta x Mt2"]))
        for metrica in ("UPT", "ATV", "ASP", "Venta x Mt2"):
            spec = grafica_kpi(d, metrica, referencia_regional(d)[metrica]).to_dict()
            filas = [v for rows in spec["datasets"].values() for v in rows if "etiqueta" in v]
            self.assertEqual(len(filas), 16 if metrica == "Venta x Mt2" else 17)
            self.assertEqual([v["valor"] for v in filas], sorted([v["valor"] for v in filas], reverse=True))
            self.assertNotIn("TOTALES", [v["codigo"] for v in filas])
            if metrica == "Venta x Mt2":
                self.assertNotIn("Z1GPP", [v["codigo"] for v in filas])
        w = load_workbook(p, data_only=True)
        for r in list(w.active.values)[3:20]:
            fila = d.set_index("codigo").loc[r[3]]
            for k, i in (("ASP", 16), ("UPT", 17), ("ATV", 18), ("% Dcto", 19)):
                self.assertAlmostEqual(fila[k], r[i])
            if r[20] > 0:
                self.assertAlmostEqual(fila["Venta x Mt2"], r[21])
        w.close()


if __name__ == "__main__":
    unittest.main()
